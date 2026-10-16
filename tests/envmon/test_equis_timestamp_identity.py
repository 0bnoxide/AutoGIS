"""Synthetic R9 imports reject ambiguous/lossy timestamps before any append."""
import dataclasses
import datetime as dt
import zipfile
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest
import xlrd

import autogis
from autogis.core.common.qa import QACollector
from autogis.core.envmon import edd_importer as importer
from autogis.core.envmon.edd_profile import LabEDDProfile
from autogis.core.envmon.equis_reader import read_equis_xls
from autogis.core.envmon.gdb_schema import compute_unique_key


PROFILE = Path(autogis.__file__).parent / "config/lab_profiles/epar4.yaml"


def _sheets(date="03/17/2025", time="09:14", *, qc=True):
    sample = dict(sys_sample_code="S1", sys_loc_code="MW1",
                  sample_matrix_code="GW", sample_date="03/11/2025",
                  sample_source="Lab" if qc else "Field",
                  sample_type_code="LCS" if qc else "N")
    test = dict(sys_sample_code="S1", lab_anl_method_name="E200.8",
                analysis_date=date, analysis_time=time,
                total_or_dissolved="T", column_number="NA",
                test_type="INITIAL", dilution_factor="1", basis="NA")
    result = {k: v for k, v in test.items()
              if k not in {"dilution_factor", "basis"}}
    result.update(chemical_name="Lead", result_value="1", result_unit="mg/L",
                  result_type_code="TRG", reportable_result="Yes",
                  test_batch_type="Analysis", test_batch_id="AB1")
    return {"EPAR4_FSample_v1": [sample], "EPAR4_TST_v1": [test],
            "EPAR4_RES_v1": [result]}


def _workbook(tmp_path, sheets, *, name="export", headers=None, iso=False):
    wb = openpyxl.Workbook()
    wb.iso_dates = iso
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        columns = list(dict.fromkeys(k for row in rows for k in row))
        ws.append([(headers or {}).get(k, k) for k in columns])
        for row in rows:
            ws.append([row.get(k) for k in columns])
    path = tmp_path / f"{name}.xlsx"
    wb.save(path)
    wb.close()
    return path


def _writer(monkeypatch):
    """Only replace the unavailable GDB effects; readers/normalizers/keys run."""
    stored = {}
    batches = []
    monkeypatch.setattr(importer, "create_or_update_gdb_schema", lambda *a: None)
    monkeypatch.setattr(importer, "create_edd_import_batch", lambda *a: "B1")

    def append(gdb, table, records, qa, batch):
        keys = stored.setdefault(table, set())
        before = len(keys)
        keys.update(compute_unique_key(dataclasses.asdict(r), table)
                    for r in records)
        return len(keys) - before, len(records) - (len(keys) - before)

    monkeypatch.setattr(importer, "append_records_idempotent", append)
    monkeypatch.setattr(importer, "finalize_batch",
                        lambda gdb, batch, qa, counts, status:
                        batches.append((status, counts)))
    monkeypatch.setattr(importer, "write_qa_to_gdb", lambda *a: None)
    return stored, batches


def _import(path, profile, qa):
    return importer.run_edd_import(path, profile, path.parent / "synthetic.gdb",
                                   "SITE", {}, {}, qa=qa)


def _keys(path, profile):
    qa = QACollector()
    rows = importer.read_edd_file(path, profile, qa)
    qc = [r for r in rows if r.get("__equis_stream") == "qc"]
    data = [r for r in rows if r.get("__equis_stream") != "qc"]
    _, results = importer.normalize_edd_rows(data, profile, "SITE", "B1", {}, {}, qa)
    results_qc = importer.normalize_qc_rows(qc, profile, "SITE", "B1", {}, qa)
    table, records = (("Env_QCResults", results_qc) if qc
                      else ("Env_AnalyticalResults", results))
    return [compute_unique_key(dataclasses.asdict(r), table) for r in records], qa


@pytest.mark.parametrize("qc", [True, False])
@pytest.mark.parametrize("date,time", [
    ("3/17/2025", "09:14"), ("03/17/2025", "9:14"),
    ("2025-03-17", "09:14"), ("02/30/2025", "09:14"),
    ("03/17/2025 09:14", ""), ("03/17/2025T09:14:00Z", ""),
    ("03/17/2025", "09:14:00"), ("03/17/2025", "09:14:17"),
    ("03/17/2025", "09:14:00.125"), ("03/17/2025", "09:14+00:00"),
    ("03/17/2025", "24:00"),
    (dt.datetime(2025, 3, 17, 0, 0, 17), "09:14"),
    (dt.date(2025, 3, 17), dt.time(9, 14, 17)),
    (dt.date(2025, 3, 17), dt.time(9, 14, 0, 125000)),
])
@pytest.mark.parametrize("side", ["EPAR4_TST_v1", "EPAR4_RES_v1"])
def test_invalid_timestamp_aborts_valid_siblings_before_any_append(
        tmp_path, monkeypatch, qc, date, time, side):
    sheets = _sheets(qc=qc)
    bad = dict(sheets[side][0], analysis_date=date, analysis_time=time)
    # The invalid TST is unmatched, and an invalid RES may lose its sample.
    # Validation must run before either join filters it away.
    bad["sys_sample_code"] = "MISSING"
    sheets[side].append(bad)
    path = _workbook(tmp_path, sheets)
    stored, batches = _writer(monkeypatch)
    qa = QACollector()
    _import(path, LabEDDProfile.load(PROFILE), qa)
    assert not stored
    assert batches[0][0] == "ERROR"
    assert all(n == 0 for n in batches[0][1].values())
    errors = [r for r in qa.records if r.category == "equis_invalid_run_timestamp"]
    assert errors and errors[0].source_sheet == side
    assert errors[0].source_row == 3


@pytest.mark.parametrize("qc", [True, False])
@pytest.mark.parametrize("time,native_time", [("09:14", dt.time(9, 14)),
                                              ("00:00", dt.time())])
def test_canonical_native_and_text_exports_keep_frozen_keys_and_reimport(
        tmp_path, monkeypatch, qc, time, native_time):
    profile = LabEDDProfile.load(PROFILE)
    text = _workbook(tmp_path, _sheets(time=time, qc=qc), name="text")
    native = _workbook(tmp_path, _sheets(dt.datetime(2025, 3, 17), native_time,
                                        qc=qc), name="native")
    keys_text, qa_text = _keys(text, profile)
    keys_native, qa_native = _keys(native, profile)
    token = f"1|INITIAL|{'' if qc else 'E200.8|'}03/17/2025@{time}"
    expected = (("SITE", "GW", "AB1", "S1", "LCS", "LEAD", "T",
                 "E200.8", token) if qc else
                ("SITE", "GW", "MW1", "S1", "2025-03-11", "LEAD", "",
                 "EQUIS_XLS", "", "T", "", token))
    assert keys_text == keys_native == [expected]
    assert not qa_text.has_blocking() and not qa_native.has_blocking()
    stored, batches = _writer(monkeypatch)
    _import(text, profile, QACollector())
    _import(native, profile, QACollector())
    table = "Env_QCResults" if qc else "Env_AnalyticalResults"
    assert stored[table] == {expected}
    assert batches[1][1]["qc_results" if qc else "analytical_results"] == 0
    assert batches[1][1]["qc_skipped" if qc else "analytical_skipped"] == 1


@pytest.mark.parametrize("date,time,token", [("", "", "1|INITIAL"),
    ("03/17/2025", "", "1|INITIAL|03/17/2025@"),
    ("", "09:14", "1|INITIAL|@09:14")])
def test_blank_discriminators_preserve_existing_degradation(tmp_path, date, time, token):
    path = _workbook(tmp_path, _sheets(date, time))
    keys, qa = _keys(path, LabEDDProfile.load(PROFILE))
    assert keys[0][-1] == token
    assert not qa.has_blocking()


@pytest.mark.parametrize("serial", [0.0, 0.5, 30.0,
    45733 + (9 * 3600 + 14 * 60 + 17) / 86400, 45733 + 0.0000000011574074])
def test_xls_loader_preserves_loss_evidence_before_legacy_conversion(
        tmp_path, monkeypatch, serial):
    sheets = _sheets()
    field = "analysis_time" if serial < 1 else "analysis_date"
    for name in ("EPAR4_TST_v1", "EPAR4_RES_v1"):
        sheets[name][0][field] = SimpleNamespace(ctype=xlrd.XL_CELL_DATE, value=serial)

    class Sheet:
        def __init__(self, rows):
            cols = list(rows[0])
            self.cells = [[SimpleNamespace(ctype=xlrd.XL_CELL_TEXT, value=c)
                           for c in cols]]
            self.cells.extend([[v if hasattr(v, "ctype") else
                                SimpleNamespace(ctype=xlrd.XL_CELL_TEXT, value=v)
                                for v in row.values()] for row in rows])
            self.nrows = len(self.cells)

        def row(self, index):
            return self.cells[index]

    book = SimpleNamespace(datemode=0, sheet_by_name=lambda n: Sheet(sheets[n]))
    monkeypatch.setattr(xlrd, "open_workbook", lambda path: book)
    stored, batches = _writer(monkeypatch)
    qa = QACollector()
    _import(tmp_path / "codec.xls", LabEDDProfile.load(PROFILE), qa)
    assert not stored
    assert batches[0][0] == "ERROR"
    assert [r for r in qa.records if r.category == "equis_invalid_run_timestamp"]


def test_header_and_source_aliases_cannot_hide_native_precision_loss(tmp_path, monkeypatch):
    sheets = _sheets(time=dt.time(9, 14, 17))
    for name in ("EPAR4_TST_v1", "EPAR4_RES_v1"):
        row = sheets[name][0]
        row["run_time"] = row.pop("analysis_time")
    path = _workbook(tmp_path, sheets, headers={"run_time": "#RUN_TIME"})
    profile = LabEDDProfile.load(PROFILE)
    profile.source_aliases["run_time"] = "analysis_time"
    stored, batches = _writer(monkeypatch)
    qa = QACollector()
    _import(path, profile, qa)
    assert not stored and batches[0][0] == "ERROR"
    assert [r for r in qa.records if r.category == "equis_invalid_run_timestamp"]


@pytest.mark.parametrize("column,native", [("analysis_date", dt.datetime(2025, 3, 17)),
                                           ("analysis_time", dt.time(9, 14))])
def test_raw_xlsx_submillisecond_precision_cannot_round_into_valid_run(
        tmp_path, monkeypatch, column, native):
    import xml.etree.ElementTree as ET
    sheets = _sheets()
    for name in ("EPAR4_TST_v1", "EPAR4_RES_v1"):
        sheets[name][0][column] = native
    path = _workbook(tmp_path, sheets)
    with zipfile.ZipFile(path) as archive:
        members = {n: archive.read(n) for n in archive.namelist()}
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    # Inject0.1ms into actual numeric timestamp source; decoder drops it.
    for member in ("xl/worksheets/sheet2.xml", "xl/worksheets/sheet3.xml"):
        root = ET.fromstring(members[member])
        row = root.find(f"{ns}sheetData/{ns}row[@r='2']")
        index = list(sheets["EPAR4_TST_v1" if member.endswith("sheet2.xml")
                            else "EPAR4_RES_v1"][0]).index(column)
        value = list(row)[index].find(f"{ns}v")
        value.text = str(Decimal(value.text) + Decimal("0.0000000011574074"))
        members[member] = ET.tostring(root)
    with zipfile.ZipFile(path, "w") as archive:
        for name, data in members.items():
            archive.writestr(name, data)
    stored, batches = _writer(monkeypatch)
    qa = QACollector()
    _import(path, LabEDDProfile.load(PROFILE), qa)
    assert not stored and batches[0][0] == "ERROR"
    assert [r for r in qa.records if r.category == "equis_invalid_run_timestamp"]


@pytest.mark.parametrize("raw", ["2025-03-17T00:00:00.0004",
    "2025-03-17T00:00:00.0000000001", "2025-03-17T00:00:00Z",
    "2025-03-17T00:00:00+00:00", "2025-03-17T00:00:00trailing"])
def test_iso_source_suffixes_cannot_escape_decoder_precision_checks(
        tmp_path, monkeypatch, raw):
    import xml.etree.ElementTree as ET
    sheets = _sheets(dt.datetime(2025, 3, 17), dt.time(9, 14))
    path = _workbook(tmp_path, sheets, iso=True)
    with zipfile.ZipFile(path) as archive:
        members = {n: archive.read(n) for n in archive.namelist()}
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    for member in ("xl/worksheets/sheet2.xml", "xl/worksheets/sheet3.xml"):
        root = ET.fromstring(members[member])
        row = root.find(f"{ns}sheetData/{ns}row[@r='2']")
        # Date is third column in the independent synthetic workbook.
        cell = list(row)[2]
        assert cell.attrib["t"] == "d"
        cell.find(f"{ns}v").text = raw
        members[member] = ET.tostring(root)
    with zipfile.ZipFile(path, "w") as archive:
        for name, data in members.items():
            archive.writestr(name, data)
    stored, batches = _writer(monkeypatch)
    qa = QACollector()
    _import(path, LabEDDProfile.load(PROFILE), qa)
    assert not stored and batches[0][0] == "ERROR"
    assert [r for r in qa.records if r.category == "equis_invalid_run_timestamp"]


def test_canonical_iso_native_cells_keep_existing_text_keys(tmp_path):
    profile = LabEDDProfile.load(PROFILE)
    path = _workbook(tmp_path, _sheets(dt.date(2025, 3, 17), dt.time(9, 14)), iso=True)
    keys, qa = _keys(path, profile)
    assert keys == [("SITE", "GW", "AB1", "S1", "LCS", "LEAD", "T",
                     "E200.8", "1|INITIAL|03/17/2025@09:14")]
    assert not qa.has_blocking()


def test_matched_bad_test_precision_cannot_be_overwritten_by_valid_result(
        tmp_path, monkeypatch):
    sheets = _sheets()
    sheets["EPAR4_TST_v1"][0]["analysis_time"] = dt.time(9, 14, 17)
    path = _workbook(tmp_path, sheets)
    stored, batches = _writer(monkeypatch)
    qa = QACollector()
    _import(path, LabEDDProfile.load(PROFILE), qa)
    assert not stored and batches[0][0] == "ERROR"
    errors = [r for r in qa.records if r.category == "equis_invalid_run_timestamp"]
    assert errors and errors[0].source_sheet == "EPAR4_TST_v1"


def test_multi_step_source_aliases_cannot_hide_precision(tmp_path, monkeypatch):
    sheets = _sheets(time=dt.time(9, 14, 17))
    for name in ("EPAR4_TST_v1", "EPAR4_RES_v1"):
        row = sheets[name][0]
        row["run_time"] = row.pop("analysis_time")
    path = _workbook(tmp_path, sheets)
    profile = LabEDDProfile.load(PROFILE)
    profile.source_aliases.update(run_time="intermediate_time", intermediate_time="analysis_time")
    stored, batches = _writer(monkeypatch)
    qa = QACollector()
    _import(path, profile, qa)
    assert not stored and batches[0][0] == "ERROR"


@pytest.mark.parametrize("qc", [True, False])
def test_canonical_runs_keep_keys_when_alone_reordered_or_subset(tmp_path, qc):
    profile = LabEDDProfile.load(PROFILE)
    first, second = _sheets(time="09:14", qc=qc), _sheets(time="16:47", qc=qc)
    together = {n: first[n] + second[n] for n in first}
    alone = _keys(_workbook(tmp_path, first, name="alone"), profile)[0][0]
    combined = _keys(_workbook(tmp_path, together, name="both"), profile)[0]
    reversed_rows = {n: list(reversed(rows)) for n, rows in together.items()}
    reverse = _keys(_workbook(tmp_path, reversed_rows, name="reverse"), profile)[0]
    assert alone == combined[0] == reverse[1]
    assert set(combined) == set(reverse) and len(set(combined)) == 2


@pytest.mark.parametrize("qc", [True, False])
def test_indistinguishable_results_and_overlength_remain_batch_blocking(tmp_path, monkeypatch, qc):
    profile = LabEDDProfile.load(PROFILE)
    sheets = _sheets(qc=qc)
    sheets["EPAR4_RES_v1"].append(dict(sheets["EPAR4_RES_v1"][0]))
    stored, batches = _writer(monkeypatch)
    qa = QACollector()
    _import(_workbook(tmp_path, sheets, name="duplicate"), profile, qa)
    assert not stored and batches[0][0] == "ERROR"
    assert [r for r in qa.records if r.category == "edd_key_collision"]
    sheets = _sheets(qc=qc)
    sheets["EPAR4_TST_v1"][0]["dilution_factor"] = "1" * 64
    qa = QACollector()
    _import(_workbook(tmp_path, sheets, name="long"), profile, qa)
    assert not stored and batches[1][0] == "ERROR"
    assert [r for r in qa.records if r.category == "edd_key_too_long"]


@pytest.mark.parametrize("datemode", [0, 1])
def test_xls_native_date_with_canonical_text_time_keeps_frozen_key(
        tmp_path, monkeypatch, datemode):
    sheets = _sheets()
    serial = xlrd.xldate.xldate_from_date_tuple((2025, 3, 17), datemode)
    for name in ("EPAR4_TST_v1", "EPAR4_RES_v1"):
        sheets[name][0]["analysis_date"] = SimpleNamespace(ctype=xlrd.XL_CELL_DATE, value=serial)

    class Sheet:
        def __init__(self, rows):
            self.headers = list(rows[0])
            self.rows = rows
            self.nrows = len(rows) + 1

        def row(self, index):
            values = self.headers if index == 0 else self.rows[index - 1].values()
            return [v if hasattr(v, "ctype") else
                    SimpleNamespace(ctype=xlrd.XL_CELL_TEXT, value=v) for v in values]

    book = SimpleNamespace(datemode=datemode, sheet_by_name=lambda n: Sheet(sheets[n]))
    monkeypatch.setattr(xlrd, "open_workbook", lambda path: book)
    keys, qa = _keys(tmp_path / "native-date.xls", LabEDDProfile.load(PROFILE))
    assert keys == [("SITE", "GW", "AB1", "S1", "LCS", "LEAD", "T",
                     "E200.8", "1|INITIAL|03/17/2025@09:14")]
    assert not qa.has_blocking()


def _rewrite_xlsx(path, mutate):
    with zipfile.ZipFile(path) as archive:
        members = {n: archive.read(n) for n in archive.namelist()}
    mutate(members)
    with zipfile.ZipFile(path, "w") as archive:
        for name, data in members.items():
            archive.writestr(name, data)


@pytest.mark.parametrize("reverse", [False, True])
@pytest.mark.parametrize("fraction", [100000, 100])
def test_three_hop_alias_precision_blocks_independent_of_mapping_order(
        tmp_path, monkeypatch, reverse, fraction):
    sheets = _sheets(time=dt.time(9, 14, 0, fraction))
    for name in ("EPAR4_TST_v1", "EPAR4_RES_v1"):
        row = sheets[name][0]
        row["raw_run"] = row.pop("analysis_time")
    path = _workbook(tmp_path, sheets)
    pairs = [("raw_run", "middle_one"), ("middle_one", "middle_two"),
             ("middle_two", "analysis_time")]
    profile = LabEDDProfile.load(PROFILE)
    profile.source_aliases.update(reversed(pairs) if reverse else pairs)
    stored, batches = _writer(monkeypatch)
    qa = QACollector()
    _import(path, profile, qa)
    assert not stored and batches[0][0] == "ERROR"
    assert [r for r in qa.records if r.category == "equis_invalid_run_timestamp"]


@pytest.mark.parametrize("native", [False, True])
@pytest.mark.parametrize("move_sheets", [False, True])
def test_relocated_opc_workbook_keeps_canonical_key(tmp_path, monkeypatch, native, move_sheets):
    import xml.etree.ElementTree as ET
    sheets = _sheets(dt.date(2025, 3, 17), dt.time(9, 14)) if native else _sheets()
    path = _workbook(tmp_path, sheets)

    def relocate(members):
        members["custom/book.xml"] = members.pop("xl/workbook.xml")
        members["custom/_rels/book.xml.rels"] = members.pop("xl/_rels/workbook.xml.rels")
        relocated_sheets = {}
        if move_sheets:
            for i in range(1, 4):
                old = f"xl/worksheets/sheet{i}.xml"
                new = f"custom/sheets/sheet{i}.xml"
                members[new] = members.pop(old)
                relocated_sheets[old] = new
            root = ET.fromstring(members["custom/_rels/book.xml.rels"])
            for rel in root:
                old = rel.attrib["Target"].lstrip("/")
                if old in relocated_sheets:
                    rel.attrib["Target"] = relocated_sheets[old].removeprefix("custom/")
            members["custom/_rels/book.xml.rels"] = ET.tostring(root)
        root = ET.fromstring(members["_rels/.rels"])
        for rel in root:
            if rel.attrib["Type"].endswith("/officeDocument"):
                rel.attrib["Target"] = "custom/book.xml"
        members["_rels/.rels"] = ET.tostring(root)
        root = ET.fromstring(members["[Content_Types].xml"])
        for item in root:
            if item.attrib.get("PartName") == "/xl/workbook.xml":
                item.attrib["PartName"] = "/custom/book.xml"
            old = item.attrib.get("PartName", "").lstrip("/")
            if old in relocated_sheets:
                item.attrib["PartName"] = "/" + relocated_sheets[old]
        members["[Content_Types].xml"] = ET.tostring(root)

    _rewrite_xlsx(path, relocate)
    stored, batches = _writer(monkeypatch)
    qa = QACollector()
    _import(path, LabEDDProfile.load(PROFILE), qa)
    assert batches[0][0] == "PASS"
    assert stored["Env_QCResults"] == {
        ("SITE", "GW", "AB1", "S1", "LCS", "LEAD", "T",
         "E200.8", "1|INITIAL|03/17/2025@09:14")}


def test_unavailable_raw_package_proof_finalizes_error_without_append(tmp_path, monkeypatch):
    # openpyxl locates the workbook through content types, whereas our proof
    # must establish its package relationship before accepting identity cells.
    path = _workbook(tmp_path, _sheets())
    _rewrite_xlsx(path, lambda members: members.pop("_rels/.rels"))
    stored, batches = _writer(monkeypatch)
    qa = QACollector()
    _import(path, LabEDDProfile.load(PROFILE), qa)
    assert not stored and batches[0][0] == "ERROR"
    assert [r for r in qa.records if r.category == "equis_invalid_run_timestamp"]


@pytest.mark.parametrize("time", ["09:14", "09:14:17", dt.time(9, 14, 17)])
def test_sample_only_run_timestamp_is_validated_without_changing_key(
        tmp_path, monkeypatch, time):
    sheets = _sheets()
    for name in ("EPAR4_TST_v1", "EPAR4_RES_v1"):
        for col in ("analysis_date", "analysis_time"):
            del sheets[name][0][col]
    sheets["EPAR4_FSample_v1"][0].update(analysis_date="03/17/2025", analysis_time=time)
    path = _workbook(tmp_path, sheets)
    stored, batches = _writer(monkeypatch)
    qa = QACollector()
    _import(path, LabEDDProfile.load(PROFILE), qa)
    if time == "09:14":
        assert batches[0][0] == "PASS"
        assert stored["Env_QCResults"] == {
            ("SITE", "GW", "AB1", "S1", "LCS", "LEAD", "T",
             "E200.8", "1|INITIAL|03/17/2025@09:14")}
    else:
        assert not stored and batches[0][0] == "ERROR"
        errors = [r for r in qa.records if r.category == "equis_invalid_run_timestamp"]
        assert errors and errors[0].source_sheet == "EPAR4_FSample_v1"


def test_sample_native_fraction_cannot_be_masked_by_empty_result_metadata(tmp_path, monkeypatch):
    import xml.etree.ElementTree as ET
    sheets = _sheets()
    for name in ("EPAR4_TST_v1", "EPAR4_RES_v1"):
        for col in ("analysis_date", "analysis_time"):
            del sheets[name][0][col]
    sample = sheets["EPAR4_FSample_v1"][0]
    sample.update(analysis_date=dt.date(2025, 3, 17), analysis_time=dt.time(9, 14))
    path = _workbook(tmp_path, sheets)

    def add_fraction(members):
        ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        root = ET.fromstring(members["xl/worksheets/sheet1.xml"])
        row = root.find(f"{ns}sheetData/{ns}row[@r='2']")
        value = list(row)[list(sample).index("analysis_time")].find(f"{ns}v")
        value.text = str(Decimal(value.text) + Decimal("0.0000000011574074"))
        members["xl/worksheets/sheet1.xml"] = ET.tostring(root)

    _rewrite_xlsx(path, add_fraction)
    stored, batches = _writer(monkeypatch)
    qa = QACollector()
    _import(path, LabEDDProfile.load(PROFILE), qa)
    assert not stored and batches[0][0] == "ERROR"
    assert [r for r in qa.records if r.category == "equis_invalid_run_timestamp"]


def test_raw_precision_proof_streams_only_identity_cells(tmp_path, monkeypatch):
    from autogis.core.envmon import equis_reader
    sheets = _sheets(dt.date(2025, 3, 17), dt.time(9, 14))
    for row in sheets["EPAR4_FSample_v1"]:
        row.update({f"unrelated_{i}": f"value-{i}" for i in range(60)})
    for i in range(14):
        sheets["EPAR4_FSample_v1"].append({
            **sheets["EPAR4_FSample_v1"][0], "sys_sample_code": f"unused-{i}"})
    path = _workbook(tmp_path, sheets)
    actual_read = zipfile.ZipFile.read

    def no_worksheet_read(archive, name, *args, **kwargs):
        assert not str(name).startswith("xl/worksheets/"), "whole worksheet bytes materialized"
        return actual_read(archive, name, *args, **kwargs)

    monkeypatch.setattr(zipfile.ZipFile, "read", no_worksheet_read)
    raw = equis_reader._xlsx_timestamp_sources(
        path, {"EPAR4_TST_v1": {"C", "D"}, "EPAR4_RES_v1": {"C", "D"}})
    assert {name: len(cells) for name, cells in raw.items()} == {
        "EPAR4_TST_v1": 2, "EPAR4_RES_v1": 2}
    stored, batches = _writer(monkeypatch)
    qa = QACollector()
    _import(path, LabEDDProfile.load(PROFILE), qa)
    assert batches[0][0] == "PASS" and len(stored["Env_QCResults"]) == 1


@pytest.mark.parametrize("date", ["03/17/2025 09:14", dt.datetime(2025, 3, 17, 0, 0, 17)])
def test_bad_sample_date_blocks_even_when_result_overrides_timestamp(tmp_path, monkeypatch, date):
    sheets = _sheets()
    sheets["EPAR4_FSample_v1"][0]["analysis_date"] = date
    path = _workbook(tmp_path, sheets)
    stored, batches = _writer(monkeypatch)
    qa = QACollector()
    _import(path, LabEDDProfile.load(PROFILE), qa)
    assert not stored and batches[0][0] == "ERROR"
    assert [r for r in qa.records if r.category == "equis_invalid_run_timestamp"]


def test_sample_reverse_alias_precision_is_blocked_before_overrides(tmp_path, monkeypatch):
    sheets = _sheets()
    sheets["EPAR4_FSample_v1"][0]["raw_run"] = dt.time(9, 14, 17)
    path = _workbook(tmp_path, sheets)
    profile = LabEDDProfile.load(PROFILE)
    profile.source_aliases.update({"middle_two": "analysis_time", "middle_one": "middle_two",
                                   "raw_run": "middle_one"})
    stored, batches = _writer(monkeypatch)
    qa = QACollector()
    _import(path, profile, qa)
    assert not stored and batches[0][0] == "ERROR"
    errors = [r for r in qa.records if r.category == "equis_invalid_run_timestamp"]
    assert errors and errors[0].source_sheet == "EPAR4_FSample_v1"


@pytest.mark.parametrize("qc", [True, False])
def test_invalid_sample_contributor_blocks_valid_sibling_batches(tmp_path, monkeypatch, qc):
    sheets = _sheets(qc=qc)
    second = _sheets(qc=qc)
    for name, rows in second.items():
        for row in rows:
            row["sys_sample_code"] = "S2"
        sheets[name].extend(rows)
    sheets["EPAR4_FSample_v1"][1]["analysis_time"] = "09:14:17"
    path = _workbook(tmp_path, sheets)
    stored, batches = _writer(monkeypatch)
    qa = QACollector()
    _import(path, LabEDDProfile.load(PROFILE), qa)
    assert not stored and batches[0][0] == "ERROR"
    assert all(n == 0 for n in batches[0][1].values())


def test_inferred_native_cell_coordinate_keeps_literal_key(tmp_path, monkeypatch):
    import xml.etree.ElementTree as ET
    path = _workbook(tmp_path, _sheets(dt.date(2025, 3, 17), dt.time(9, 14)))

    def remove_coordinate(members):
        ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        root = ET.fromstring(members["xl/worksheets/sheet3.xml"])
        row = root.find(f"{ns}sheetData/{ns}row[@r='2']")
        list(row)[3].attrib.pop("r")
        members["xl/worksheets/sheet3.xml"] = ET.tostring(root)

    _rewrite_xlsx(path, remove_coordinate)
    stored, batches = _writer(monkeypatch)
    qa = QACollector()
    _import(path, LabEDDProfile.load(PROFILE), qa)
    assert stored["Env_QCResults"] == {
        ("SITE", "GW", "AB1", "S1", "LCS", "LEAD", "T", "E200.8", "1|INITIAL|03/17/2025@09:14")}
    assert batches[0][0] == "PASS" and not qa.has_blocking()


def _timestamp_alias_pairs(role, shape):
    if shape == "one_hop":
        return [("raw_run", role)]
    count = 6 if shape == "reverse_six" else 3
    names = ["raw_run"] + [f"middle_{i}" for i in range(1, count)] + [role]
    pairs = list(zip(names, names[1:]))
    if shape.startswith("cycle"):
        pairs.append((role, "raw_run"))
    return list(reversed(pairs)) if shape.startswith("reverse") or shape == "cycle_reverse" else pairs


@pytest.mark.parametrize("role,bad", [
    ("analysis_time", "09:14:17"), ("analysis_time", "9:14"),
    ("analysis_time", "09:14:00.0001"), ("analysis_time", "09:14Z"),
    ("analysis_date", "3/17/2025"), ("analysis_date", "2025-03-17"),
    ("analysis_date", "03/17/2025 09:14"), ("analysis_date", "02/30/2025")])
@pytest.mark.parametrize("side", ["EPAR4_FSample_v1", "EPAR4_TST_v1", "EPAR4_RES_v1"])
def test_reverse_alias_raw_text_is_validated_before_values_disappear(
        tmp_path, monkeypatch, role, bad, side):
    sheets = _sheets(qc=side != "EPAR4_RES_v1")
    for row in sheets[side]:
        row.pop(role, None)
        row["raw_run"] = bad
    profile = LabEDDProfile.load(PROFILE)
    profile.source_aliases.update(_timestamp_alias_pairs(role, "reverse_three"))
    path = _workbook(tmp_path, sheets)
    stored, batches = _writer(monkeypatch)
    qa = QACollector()
    _import(path, profile, qa)
    assert not stored and batches[0][0] == "ERROR"
    errors = [r for r in qa.records if r.category == "equis_invalid_run_timestamp"]
    assert errors and errors[0].source_sheet == side


@pytest.mark.parametrize("role", ["analysis_date", "analysis_time"])
@pytest.mark.parametrize("shape", ["one_hop", "forward_three", "reverse_three",
                                    "reverse_six", "cycle_forward", "cycle_reverse"])
@pytest.mark.parametrize("invalid", [True, False])
def test_text_alias_role_reachability_preserves_existing_canonical_keys(
        tmp_path, monkeypatch, role, shape, invalid):
    sheets = _sheets()
    value = "03/17/2025" if role == "analysis_date" else "09:14"
    if invalid:
        value = "3/17/2025" if role == "analysis_date" else "09:14:17"
    for name in ("EPAR4_TST_v1", "EPAR4_RES_v1"):
        row = sheets[name][0]
        row["raw_run"] = value
        del row[role]
    profile = LabEDDProfile.load(PROFILE)
    profile.source_aliases.update(_timestamp_alias_pairs(role, shape))
    path = _workbook(tmp_path, sheets)
    stored, batches = _writer(monkeypatch)
    qa = QACollector()
    _import(path, profile, qa)
    if invalid:
        assert not stored and batches[0][0] == "ERROR"
        assert [r for r in qa.records if r.category == "equis_invalid_run_timestamp"]
    else:
        # The real value-alias pass stays single-pass. Preserve its exact
        # existing missing-field degradation for reverse chains and cycles.
        token = "1|INITIAL|03/17/2025@09:14"
        if shape in {"reverse_three", "reverse_six", "cycle_reverse"}:
            token = "1|INITIAL|@09:14" if role == "analysis_date" else "1|INITIAL|03/17/2025@"
        assert batches[0][0] == "PASS"
        assert stored["Env_QCResults"] == {
            ("SITE", "GW", "AB1", "S1", "LCS", "LEAD", "T", "E200.8", token)}


@pytest.mark.parametrize("raw,blocking", [("", False), ("09:14", False), ("09:14:17", True)])
def test_alias_raw_text_is_validated_even_when_canonical_value_overrides_it(
        tmp_path, monkeypatch, raw, blocking):
    sheets = _sheets(time="16:47")
    for name in ("EPAR4_TST_v1", "EPAR4_RES_v1"):
        sheets[name][0]["raw_run"] = raw
    profile = LabEDDProfile.load(PROFILE)
    profile.source_aliases["raw_run"] = "analysis_time"
    stored, batches = _writer(monkeypatch)
    qa = QACollector()
    _import(_workbook(tmp_path, sheets), profile, qa)
    if blocking:
        assert not stored and batches[0][0] == "ERROR"
    else:
        assert batches[0][0] == "PASS"
        assert stored["Env_QCResults"] == {
            ("SITE", "GW", "AB1", "S1", "LCS", "LEAD", "T",
             "E200.8", "1|INITIAL|03/17/2025@16:47")}


@pytest.mark.parametrize("role,token", [("analysis_date", "1|INITIAL|@09:14"),
                                       ("analysis_time", "1|INITIAL|03/17/2025@")])
def test_reverse_alias_blank_contributor_keeps_existing_degraded_key(
        tmp_path, monkeypatch, role, token):
    sheets = _sheets()
    for name in ("EPAR4_TST_v1", "EPAR4_RES_v1"):
        row = sheets[name][0]
        row["raw_run"] = ""
        del row[role]
    profile = LabEDDProfile.load(PROFILE)
    profile.source_aliases.update(_timestamp_alias_pairs(role, "reverse_three"))
    stored, batches = _writer(monkeypatch)
    qa = QACollector()
    _import(_workbook(tmp_path, sheets), profile, qa)
    assert batches[0][0] == "PASS" and not qa.has_blocking()
    assert stored["Env_QCResults"] == {
        ("SITE", "GW", "AB1", "S1", "LCS", "LEAD", "T", "E200.8", token)}


def _formula_workbook(tmp_path, role, cache, *, side=None, alias=False, qc=True,
                      midnight=False, shared=False, sample_only=False):
    """Write explicit raw cached/uncached formulas without evaluating them."""
    import xml.etree.ElementTree as ET
    date, time = dt.date(2025, 3, 17), dt.time(0 if midnight else 9, 0 if midnight else 14)
    sheets = _sheets(date, time, qc=qc)
    if sample_only:
        for name in ("EPAR4_TST_v1", "EPAR4_RES_v1"):
            sheets[name][0].pop(role)
    sides = [side] if side else ["EPAR4_TST_v1", "EPAR4_RES_v1"]
    source = "raw_run" if alias else role
    for name in sides:
        if alias:
            sheets[name][0].pop(role, None)
        sheets[name][0][source] = date if role == "analysis_date" else time
    profile = LabEDDProfile.load(PROFILE)
    if alias:
        profile.source_aliases.update(_timestamp_alias_pairs(role, "reverse_three"))
    path = _workbook(tmp_path, sheets)

    def mutate(members):
        ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        for name in sides:
            part = f"xl/worksheets/sheet{list(sheets).index(name) + 1}.xml"
            root = ET.fromstring(members[part])
            column = openpyxl.utils.get_column_letter(list(sheets[name][0]).index(source) + 1)
            cell = root.find(f".//{ns}c[@r='{column}2']")
            if cache != "blank":
                formula = ET.SubElement(cell, ns + "f")
                if shared:
                    formula.attrib.update(t="shared", si="0")
                else:
                    formula.text = "NOW()" if role == "analysis_date" else "TIME(9,14,0)"
            value = cell.find(ns + "v")
            if cache in {"absent", "blank"}:
                cell.remove(value)
            elif cache in {"empty", "emptytext", "whitespace"}:
                if cache != "empty":
                    cell.attrib["t"] = "str"
                value.text = "   " if cache == "whitespace" else None
            elif cache == "text":
                cell.attrib["t"] = "str"
                value.text = "03/17/2025" if role == "analysis_date" else "00:00" if midnight else "09:14"
            members[part] = ET.tostring(root)

    _rewrite_xlsx(path, mutate)
    return path, profile


@pytest.mark.parametrize("role", ["analysis_date", "analysis_time"])
@pytest.mark.parametrize("side", ["EPAR4_FSample_v1", "EPAR4_TST_v1", "EPAR4_RES_v1"])
@pytest.mark.parametrize("cache", ["absent", "empty", "emptytext", "whitespace"])
@pytest.mark.parametrize("alias", [False, True])
def test_formula_without_usable_cache_blocks_all_run_contributors(
        tmp_path, monkeypatch, role, side, cache, alias):
    path, profile = _formula_workbook(tmp_path, role, cache, side=side, alias=alias,
                                      qc=side != "EPAR4_RES_v1", shared=cache == "empty")
    stored, batches = _writer(monkeypatch)
    qa = QACollector()
    _import(path, profile, qa)
    assert not stored and batches[0][0] == "ERROR"
    assert all(value == 0 for value in batches[0][1].values())
    assert any(r.category == "equis_invalid_run_timestamp" and r.source_sheet == side
               for r in qa.records)


@pytest.mark.parametrize("role", ["analysis_date", "analysis_time"])
@pytest.mark.parametrize("cache", ["native", "text", "blank"])
@pytest.mark.parametrize("qc", [True, False])
@pytest.mark.parametrize("midnight", [False, True])
def test_verified_formula_caches_and_true_blanks_keep_literal_keys(
        tmp_path, monkeypatch, role, cache, qc, midnight):
    path, profile = _formula_workbook(tmp_path, role, cache, qc=qc, midnight=midnight)
    stored, batches = _writer(monkeypatch)
    qa = QACollector()
    _import(path, profile, qa)
    date, time = "03/17/2025", "00:00" if midnight else "09:14"
    if cache == "blank":
        if role == "analysis_date":
            date = ""
        else:
            time = ""
    token = f"1|INITIAL|{date}@{time}"
    key = (("SITE", "GW", "AB1", "S1", "LCS", "LEAD", "T", "E200.8", token)
           if qc else ("SITE", "GW", "MW1", "S1", "2025-03-11", "LEAD", "",
                       "EQUIS_XLS", "", "T", "", f"1|INITIAL|E200.8|{date}@{time}"))
    assert stored["Env_QCResults" if qc else "Env_AnalyticalResults"] == {key}
    assert batches[0][0] == "PASS" and not qa.has_blocking()


@pytest.mark.parametrize("role", ["analysis_date", "analysis_time"])
@pytest.mark.parametrize("cache", ["native", "text", "blank"])
def test_verified_reverse_alias_formula_cache_preserves_missing_role_key(
        tmp_path, monkeypatch, role, cache):
    path, profile = _formula_workbook(tmp_path, role, cache, alias=True)
    stored, batches = _writer(monkeypatch)
    qa = QACollector()
    _import(path, profile, qa)
    token = "1|INITIAL|@09:14" if role == "analysis_date" else "1|INITIAL|03/17/2025@"
    assert stored["Env_QCResults"] == {
        ("SITE", "GW", "AB1", "S1", "LCS", "LEAD", "T", "E200.8", token)}
    assert batches[0][0] == "PASS" and not qa.has_blocking()


@pytest.mark.parametrize("role", ["analysis_date", "analysis_time"])
@pytest.mark.parametrize("cache", ["native", "text", "blank"])
def test_verified_sample_only_formula_cache_keeps_inherited_key(tmp_path, monkeypatch, role, cache):
    path, profile = _formula_workbook(tmp_path, role, cache, side="EPAR4_FSample_v1", sample_only=True)
    stored, batches = _writer(monkeypatch)
    qa = QACollector()
    _import(path, profile, qa)
    token = "1|INITIAL|03/17/2025@09:14"
    if cache == "blank":
        token = "1|INITIAL|@09:14" if role == "analysis_date" else "1|INITIAL|03/17/2025@"
    assert stored["Env_QCResults"] == {
        ("SITE", "GW", "AB1", "S1", "LCS", "LEAD", "T", "E200.8", token)}
    assert batches[0][0] == "PASS" and not qa.has_blocking()


@pytest.mark.parametrize("bad_first", [True, False])
def test_uncached_formula_batch_preserves_valid_siblings_in_both_orders(tmp_path, bad_first):
    from autogis.core.envmon.batch_workbook_importer import run_batch_import
    bad, _ = _formula_workbook(tmp_path, "analysis_time", "absent", qc=False)
    good_sheets = _sheets(qc=False)
    for rows in good_sheets.values():
        rows[0]["sys_sample_code"] = "GOOD"
    good = _workbook(tmp_path, good_sheets, name="good")
    paths = [bad, good] if bad_first else [good, bad]
    result = run_batch_import([dict(workbook_path=str(p), profile_path=str(PROFILE), site_id="SITE")
                               for p in paths])
    assert {Path(r.WorkbookPath).name: r.Status for r in result.manifest} == {
        bad.name: "ERROR", good.name: "OK"}
    assert [r.SampleID for r in result.all_samples] == ["GOOD"]
    assert [r.SampleID for r in result.all_results] == ["GOOD"]


@pytest.mark.parametrize("cache", ["absent", "empty", "whitespace", "native"])
@pytest.mark.parametrize("mode", ["api", "manifest", "directory"])
def test_formula_cache_guard_reaches_actual_batch_api_and_cli(tmp_path, monkeypatch, cache, mode):
    import csv
    from click.testing import CliRunner
    from autogis.adapters.cli import autogis as command
    from autogis.core.envmon.batch_workbook_importer import run_batch_import
    monkeypatch.setenv("AUTOGIS_RUN_HISTORY", "off")
    inputs = tmp_path / "inputs"
    inputs.mkdir()
    path, _ = _formula_workbook(inputs, "analysis_time", cache, qc=False)
    rows = [dict(workbook_path=str(path), profile_path=str(PROFILE), site_id="SITE")]
    if mode == "api":
        result = run_batch_import(rows)
        assert result.manifest[0].Status == ("OK" if cache == "native" else "ERROR")
        assert len(result.all_results) == (1 if cache == "native" else 0)
        assert len(result.all_samples) == (1 if cache == "native" else 0)
        return
    manifest = tmp_path / "manifest.csv"
    with manifest.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    options = (["--manifest", str(manifest)] if mode == "manifest" else
               ["--edd-dir", str(inputs), "--profile", str(PROFILE), "--site", "SITE"])
    out = tmp_path / "out"
    invoked = CliRunner().invoke(command, ["envmon", "batch-import-workbooks", *options,
                                "--output-dir", str(out)], catch_exceptions=False)
    assert invoked.exit_code == (0 if cache == "native" else 1)
    assert (out / "result_records.csv").exists() == (cache == "native")
    assert (out / "sample_records.csv").exists() == (cache == "native")


def _coordinate_formula_workbook(tmp_path, shape, role="analysis_time", cache="absent",
                                 *, side=None, alias=False, midnight=False):
    import xml.etree.ElementTree as ET
    path, profile = _formula_workbook(tmp_path, role, cache, side=side, alias=alias,
                                      qc=False, midnight=midnight)
    names = [side] if side else ["EPAR4_TST_v1", "EPAR4_RES_v1"]
    source = "raw_run" if alias else role
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    targets = {name: list(next(wb[name].iter_rows(values_only=True))).index(source) + 1
               for name in names}
    wb.close()

    def mutate(members):
        ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        sheet_order = list(_sheets())
        for name, column in targets.items():
            part = f"xl/worksheets/sheet{sheet_order.index(name) + 1}.xml"
            root = ET.fromstring(members[part])
            rows = root.findall(f"{ns}sheetData/{ns}row")
            data = rows[1]
            target = data.find(f"{ns}c[@r='{openpyxl.utils.get_column_letter(column)}2']")
            if shape == "gap":
                for cell in root.iter(ns + "c"):
                    col, row = openpyxl.utils.cell.coordinate_from_string(cell.attrib["r"])
                    index = openpyxl.utils.column_index_from_string(col)
                    cell.attrib["r"] = f"{openpyxl.utils.get_column_letter(index + (index > 1))}{row}"
                unused = ET.Element(ns + "c", r="B1", t="inlineStr")
                ET.SubElement(ET.SubElement(unused, ns + "is"), ns + "t").text = "unused"
                rows[0].insert(1, unused)
                dimension = root.find(ns + "dimension")
                bounds = openpyxl.utils.range_boundaries(dimension.attrib["ref"])
                dimension.attrib["ref"] = f"A1:{openpyxl.utils.get_column_letter(bounds[2] + 1)}2"
                targets[name] += 1
            if shape in {"row", "both", "all", "gap"}:
                data.attrib.pop("r")
            if shape == "all":
                rows[0].attrib.pop("r")
                for cell in root.iter(ns + "c"):
                    cell.attrib.pop("r")
            elif shape == "mixed":
                for index, cell in enumerate(data):
                    if index % 2 == 0 or cell is target:
                        cell.attrib.pop("r", None)
            elif shape != "row":
                target.attrib.pop("r")
            if shape == "noncanonical_cell":
                target.attrib["r"] = f"{openpyxl.utils.get_column_letter(column)}02"
            elif shape == "inconsistent_cell":
                target.attrib["r"] = f"{openpyxl.utils.get_column_letter(column)}3"
            elif shape == "noncanonical_row":
                data.attrib["r"] = "02"
            members[part] = ET.tostring(root)

    _rewrite_xlsx(path, mutate)
    return path, profile, targets


_COORDINATE_CASES = (
    [(shape, "analysis_time", "absent", None, False, False)
     for shape in ("cell", "row", "both", "mixed", "all", "gap")]
    + [(shape, "analysis_date", "absent", None, False, False)
       for shape in ("both", "all", "gap")]
    + [("cell", "analysis_time", cache, None, False, False) for cache in ("empty", "whitespace")]
    + [("both", role, "absent", side, True, False) for role, side in (
        ("analysis_date", "EPAR4_FSample_v1"), ("analysis_time", "EPAR4_FSample_v1"),
        ("analysis_time", "EPAR4_TST_v1"), ("analysis_time", "EPAR4_RES_v1"))]
    + [(shape, "analysis_time", "native", None, False, False)
       for shape in ("cell", "row", "both", "mixed", "all", "gap")]
    + [("both", role, cache, None, False, False) for role, cache in (
        ("analysis_date", "text"), ("analysis_time", "text"),
        ("analysis_date", "blank"), ("analysis_time", "blank"))]
    + [("all", "analysis_time", cache, None, True, False) for cache in ("native", "text", "blank")]
    + [("both", "analysis_date", "native", "EPAR4_FSample_v1", False, False)]
    + [(shape, "analysis_time", "native", None, False, True) for shape in ("all", "gap")]
    + [(shape, "analysis_time", "absent", None, False, False)
       for shape in ("noncanonical_cell", "inconsistent_cell", "noncanonical_row")]
)


@pytest.mark.parametrize("shape,role,cache,side,alias,midnight", _COORDINATE_CASES)
def test_timestamp_proof_tracks_declared_and_inferred_positions(
        tmp_path, monkeypatch, shape, role, cache, side, alias, midnight):
    import json
    from autogis.core.envmon import equis_reader
    path, profile, targets = _coordinate_formula_workbook(
        tmp_path, shape, role, cache, side=side, alias=alias, midnight=midnight)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    decoded = {name: dict(requested_coordinate=f"{openpyxl.utils.get_column_letter(col)}2",
                          value=repr(wb[name].cell(2, col).value)) for name, col in targets.items()}
    wb.close()
    proof = equis_reader._xlsx_timestamp_sources(path, {
        name: {openpyxl.utils.get_column_letter(col)} for name, col in targets.items()})
    qa = QACollector()
    reader = importer.read_edd_file(path, profile, qa)
    stored, batches = _writer(monkeypatch)
    _import(path, profile, qa)
    (tmp_path / "observed.json").write_text(json.dumps(dict(
        decoded=decoded, proof=proof, reader_keys=[r.get("__equis_method_dilution_key") for r in reader],
        qa=[dataclasses.asdict(r) for r in qa.records], status=batches[0][0],
        stored={k: list(v) for k, v in stored.items()}), indent=2, default=str))
    if cache in {"absent", "empty", "whitespace"}:
        assert not stored and batches[0][0] == "ERROR"
        assert any(r.category == "equis_invalid_run_timestamp" for r in qa.records)
    else:
        date, time = "03/17/2025", "00:00" if midnight else "09:14"
        if (cache == "blank" or alias) and side is None:
            if role == "analysis_date":
                date = ""
            else:
                time = ""
        assert stored["Env_AnalyticalResults"] == {(
            "SITE", "GW", "MW1", "S1", "2025-03-11", "LEAD", "", "EQUIS_XLS", "", "T", "",
            f"1|INITIAL|E200.8|{date}@{time}")}
        assert batches[0][0] == "PASS" and not qa.has_blocking()


@pytest.mark.parametrize("mode", ["api", "manifest", "directory"])
@pytest.mark.parametrize("cache", ["absent", "native"])
def test_inferred_formula_proof_reaches_batch_api_and_cli(tmp_path, monkeypatch, mode, cache):
    import csv
    from click.testing import CliRunner
    from autogis.adapters.cli import autogis as command
    from autogis.core.envmon.batch_workbook_importer import run_batch_import
    monkeypatch.setenv("AUTOGIS_RUN_HISTORY", "off")
    inputs = tmp_path / "inputs"
    inputs.mkdir()
    path, _, _ = _coordinate_formula_workbook(inputs, "all", cache=cache)
    rows = [dict(workbook_path=str(path), profile_path=str(PROFILE), site_id="SITE")]
    if mode == "api":
        result = run_batch_import(rows)
        assert result.manifest[0].Status == ("OK" if cache == "native" else "ERROR")
        assert len(result.all_results) == (1 if cache == "native" else 0)
        assert len(result.all_samples) == (1 if cache == "native" else 0)
        return
    manifest = tmp_path / "manifest.csv"
    with manifest.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    options = (["--manifest", str(manifest)] if mode == "manifest" else
               ["--edd-dir", str(inputs), "--profile", str(PROFILE), "--site", "SITE"])
    out = tmp_path / "out"
    invoked = CliRunner().invoke(command, ["envmon", "batch-import-workbooks", *options,
                                "--output-dir", str(out)], catch_exceptions=False)
    assert invoked.exit_code == (0 if cache == "native" else 1)
    assert (out / "result_records.csv").exists() == (cache == "native")
    if cache == "native":
        with (out / "result_records.csv").open(newline="", encoding="utf-8-sig") as stream:
            assert [r["MethodDilutionKey"] for r in csv.DictReader(stream)] == [
                "1|INITIAL|E200.8|03/17/2025@09:14"]
    else:
        assert not (out / "sample_records.csv").exists()


@pytest.mark.parametrize("bad_first", [True, False])
def test_inferred_uncached_formula_batch_keeps_valid_siblings(tmp_path, bad_first):
    from autogis.core.envmon.batch_workbook_importer import run_batch_import
    bad, _, _ = _coordinate_formula_workbook(tmp_path, "gap")
    good_sheets = _sheets(qc=False)
    for rows in good_sheets.values():
        rows[0]["sys_sample_code"] = "GOOD"
    good = _workbook(tmp_path, good_sheets, name="good")
    result = run_batch_import([dict(workbook_path=str(p), profile_path=str(PROFILE), site_id="SITE")
                               for p in ([bad, good] if bad_first else [good, bad])])
    assert {Path(r.WorkbookPath).name: r.Status for r in result.manifest} == {
        bad.name: "ERROR", good.name: "OK"}
    assert [r.SampleID for r in result.all_results] == ["GOOD"]
    assert [r.SampleID for r in result.all_samples] == ["GOOD"]
