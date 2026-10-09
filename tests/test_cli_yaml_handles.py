"""CLI YAML inputs are read-and-closed, not left to the garbage collector.

A bare ``open()`` inside ``yaml.safe_load(...)`` leaks the handle until GC:
a ResourceWarning in the suite, and a lingering lock on Windows for as long
as the process lives. ``build-survey-form`` is the command the suite flagged.
"""
import gc
import warnings

import pytest
from click.testing import CliRunner

from autogis.adapters.cli import autogis


def test_build_survey_form_closes_its_yaml_inputs(tmp_path):
    cfg = tmp_path / "site.yaml"
    cfg.write_text("site_id: S1\n", encoding="utf-8")
    empty = tmp_path / "empty.yaml"
    empty.write_text("{}\n", encoding="utf-8")

    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always", ResourceWarning)
        CliRunner().invoke(autogis, [
            "envmon", "build-survey-form",
            "--site", str(cfg), "--analytes", str(empty),
            "--event", str(empty), "--out", str(tmp_path / "form.xlsx"),
        ])
        gc.collect()

    leaked = [w for w in caught if issubclass(w.category, ResourceWarning)]
    assert not leaked, [str(w.message) for w in leaked]


def _bad_inputs(tmp_path):
    """Malformed YAML and a directory: both pass ``click.Path(exists=True)``."""
    bad = tmp_path / "bad.yaml"
    bad.write_text("site_id: [unclosed\n", encoding="utf-8")
    a_dir = tmp_path / "a dir é"
    a_dir.mkdir()
    return bad, a_dir


def test_build_survey_form_bad_yaml_is_a_clean_error(tmp_path):
    empty = tmp_path / "empty.yaml"
    empty.write_text("{}\n", encoding="utf-8")
    for bad in _bad_inputs(tmp_path):
        result = CliRunner().invoke(autogis, [
            "envmon", "build-survey-form",
            "--site", str(bad), "--analytes", str(empty),
            "--event", str(empty), "--out", str(tmp_path / "form.xlsx"),
        ])
        assert result.exit_code == 1, result.output
        assert result.exception is None or isinstance(result.exception,
                                                      SystemExit)
        assert "cannot read" in result.output and str(bad) in result.output


def test_validate_survey_form_bad_yaml_is_a_clean_error(tmp_path):
    from autogis.core.envmon.survey123_form_builder import build_xlsform
    form = tmp_path / "form.xlsx"
    build_xlsform({"site_id": "S1"}, {}, {}).save(form)
    for bad in _bad_inputs(tmp_path):
        result = CliRunner().invoke(autogis, [
            "envmon", "validate-survey-form", str(form),
            "--site-config", str(bad),
        ])
        assert result.exit_code == 1, result.output
        assert result.exception is None or isinstance(result.exception,
                                                      SystemExit)
        assert "cannot read" in result.output and str(bad) in result.output


def _build(tmp_path, site, analytes, event, out=None):
    return CliRunner().invoke(autogis, [
        "envmon", "build-survey-form",
        "--site", str(site), "--analytes", str(analytes),
        "--event", str(event), "--out", str(out or tmp_path / "form.xlsx"),
    ])


def test_build_survey_form_rejects_empty_and_non_mapping_yaml(tmp_path):
    """#549: an empty or list-shaped required input must fail loudly. Treating
    an empty file as {} would write a form with no site, wells or analytes
    and exit 0."""
    obj = tmp_path / "obj.yaml"
    obj.write_text("{}\n", encoding="utf-8")
    empty = tmp_path / "empty.yaml"
    empty.write_text("", encoding="utf-8")
    a_list = tmp_path / "list.yaml"
    a_list.write_text("- 1\n- 2\n", encoding="utf-8")
    for args, bad in [((empty, obj, obj), empty), ((a_list, obj, obj), a_list),
                      ((obj, a_list, obj), a_list), ((obj, obj, empty), empty)]:
        result = _build(tmp_path, *args)
        assert result.exit_code == 1, result.output
        assert isinstance(result.exception, SystemExit), result.exception
        assert "expected a YAML mapping" in result.output
        assert str(bad) in result.output
        assert not (tmp_path / "form.xlsx").exists()


def test_build_survey_form_out_directory_is_a_usage_error(tmp_path):
    obj = tmp_path / "obj.yaml"
    obj.write_text("{}\n", encoding="utf-8")
    out_dir = tmp_path / "out dir"
    out_dir.mkdir()
    result = _build(tmp_path, obj, obj, obj, out=out_dir)
    assert result.exit_code == 2, result.output
    assert isinstance(result.exception, SystemExit), result.exception
    assert "is a directory" in result.output


@pytest.mark.parametrize("target_exists", [True, False], ids=["live", "dangling"])
def test_build_survey_form_rejects_symlink_output(tmp_path, monkeypatch, target_exists):
    """#560: reject the output link before writing a stage or history record."""
    import os

    obj = tmp_path / "obj.yaml"
    obj.write_text("{}\n", encoding="utf-8")
    target = tmp_path / "target.xlsx"
    if target_exists:
        target.write_bytes(b"previous form")
    out = tmp_path / "form.xlsx"
    try:
        out.symlink_to(target)
    except OSError as exc:
        pytest.skip(f"symlink creation unavailable: {exc}")
    original_link = out.lstat()
    original_link_text = os.readlink(out)
    original_target = target.stat() if target_exists else None
    history = tmp_path / "history.csv"
    monkeypatch.setenv("AUTOGIS_RUN_HISTORY", str(history))

    result = _build(tmp_path, obj, obj, obj, out)
    assert result.exit_code == 2, result.output
    assert isinstance(result.exception, SystemExit), result.exception
    assert "--out" in result.output and "symbolic link" in result.output
    assert str(out) in result.output
    assert "XLSForm written" not in result.output
    assert out.is_symlink() and os.readlink(out) == original_link_text
    link_stat = out.lstat()
    assert (link_stat.st_ino, link_stat.st_mode, link_stat.st_mtime_ns) == (
        original_link.st_ino, original_link.st_mode, original_link.st_mtime_ns)
    if target_exists:
        assert target.read_bytes() == b"previous form"
        target_stat = target.stat()
        assert (target_stat.st_ino, target_stat.st_mode, target_stat.st_mtime_ns) == (
            original_target.st_ino, original_target.st_mode, original_target.st_mtime_ns)
    else:
        assert not target.exists()
    assert list(tmp_path.glob("*.tmp")) == []
    assert not history.exists()


@pytest.mark.parametrize("kind", ["hardlink", "fifo", "socket"])
def test_build_survey_form_rejects_unsafe_destination(tmp_path, monkeypatch, kind):
    """#561: publication cannot detach aliases or replace special files."""
    import os
    import socket

    if kind != "hardlink" and os.name != "posix":
        pytest.skip("POSIX special files")
    obj = tmp_path / "obj.yaml"
    obj.write_text("{}\n", encoding="utf-8")
    out = tmp_path / "form.xlsx"
    alias = tmp_path / "consumer.xlsx"
    connection = None
    if kind == "hardlink":
        out.write_bytes(b"previous form")
        os.link(out, alias)
    elif kind == "fifo":
        os.mkfifo(out)
    else:
        connection = socket.socket(socket.AF_UNIX)
        connection.bind(str(out))
    before = out.lstat()
    history = tmp_path / "history.csv"
    monkeypatch.setenv("AUTOGIS_RUN_HISTORY", str(history))
    try:
        result = _build(tmp_path, obj, obj, obj, out)
        assert result.exit_code == 2, result.output
        assert isinstance(result.exception, SystemExit), result.exception
        assert "single-link regular file" in result.output and str(out) in result.output
        assert "XLSForm written" not in result.output
        after = out.lstat()
        assert (after.st_ino, after.st_mode, after.st_nlink, after.st_mtime_ns) == (
            before.st_ino, before.st_mode, before.st_nlink, before.st_mtime_ns)
        if kind == "hardlink":
            assert out.read_bytes() == alias.read_bytes() == b"previous form"
            assert os.path.samefile(out, alias)
        assert list(tmp_path.glob("*.tmp")) == []
        assert not history.exists()
    finally:
        if connection:
            connection.close()


def test_build_survey_form_output_inspection_error_is_clean(tmp_path, monkeypatch):
    from autogis.adapters import cli

    obj = tmp_path / "obj.yaml"
    obj.write_text("{}\n", encoding="utf-8")
    out = tmp_path / "form.xlsx"
    real_lstat = cli.Path.lstat

    def denied_lstat(path, *args, **kwargs):
        if path == out:
            raise PermissionError("output inspection denied")
        return real_lstat(path, *args, **kwargs)

    monkeypatch.setattr(cli.Path, "lstat", denied_lstat)
    history = tmp_path / "history.csv"
    monkeypatch.setenv("AUTOGIS_RUN_HISTORY", str(history))
    result = _build(tmp_path, obj, obj, obj, out)
    assert result.exit_code == 2, result.output
    assert "cannot inspect output" in result.output and "output inspection denied" in result.output
    assert str(out) in result.output
    assert not out.exists() and not history.exists()
    assert list(tmp_path.glob("*.tmp")) == []


def test_validate_survey_form_empty_config_still_means_not_supplied(tmp_path):
    from autogis.core.envmon.survey123_form_builder import build_xlsform
    form = tmp_path / "form.xlsx"
    build_xlsform({"site_id": "S1"}, {}, {}).save(form)
    empty = tmp_path / "empty.yaml"
    empty.write_text("", encoding="utf-8")
    a_list = tmp_path / "list.yaml"
    a_list.write_text("- 1\n", encoding="utf-8")

    ok = CliRunner().invoke(autogis, [
        "envmon", "validate-survey-form", str(form), "--event-config", str(empty)])
    assert not isinstance(ok.exception, AttributeError), ok.exception

    for opt in ("--site-config", "--event-config", "--analyte-dict"):
        result = CliRunner().invoke(autogis, [
            "envmon", "validate-survey-form", str(form), opt, str(a_list)])
        assert result.exit_code == 1, result.output
        assert isinstance(result.exception, SystemExit), result.exception
        assert "expected a YAML mapping" in result.output


def test_build_survey_form_creates_missing_out_directory(tmp_path):
    """#550 (owner decision: option B): an --out in a missing directory
    creates that directory instead of escaping as a raw FileNotFoundError."""
    obj = tmp_path / "obj.yaml"
    obj.write_text("{}\n", encoding="utf-8")
    out = tmp_path / "new dir é" / "nested" / "form.xlsx"
    result = _build(tmp_path, obj, obj, obj, out=out)
    assert result.exit_code == 0, result.output
    assert out.is_file()


def test_build_survey_form_unwritable_out_is_a_clean_error(tmp_path):
    """Other write failures (here: a parent that is a regular file) still
    surface as a clean error, not a traceback."""
    obj = tmp_path / "obj.yaml"
    obj.write_text("{}\n", encoding="utf-8")
    a_file = tmp_path / "a_file"
    a_file.write_text("", encoding="utf-8")
    out = a_file / "form.xlsx"
    result = _build(tmp_path, obj, obj, obj, out=out)
    assert result.exit_code == 1, result.output
    assert isinstance(result.exception, SystemExit), result.exception
    assert "cannot write" in result.output and str(out) in result.output


def test_build_survey_form_failed_save_keeps_existing_form(tmp_path, monkeypatch):
    """#551: a save that dies partway (e.g. disk full) must not truncate an
    existing --out form or leave a partial file behind."""
    from openpyxl import Workbook

    obj = tmp_path / "obj.yaml"
    obj.write_text("{}\n", encoding="utf-8")
    out = tmp_path / "form.xlsx"
    out.write_bytes(b"previous form")

    def partial_save(self, filename):
        filename.write(b"PK\x03\x04trunc")
        raise OSError(28, "No space left on device")

    monkeypatch.setattr(Workbook, "save", partial_save)
    result = _build(tmp_path, obj, obj, obj, out=out)
    assert result.exit_code == 1, result.output
    assert "cannot write" in result.output
    assert out.read_bytes() == b"previous form"
    assert sorted(p.name for p in tmp_path.iterdir()) == ["form.xlsx", "obj.yaml"]


def test_build_survey_form_failed_save_reports_real_error(tmp_path, monkeypatch):
    """#551: a cleanup failure must not mask the real openpyxl save error."""
    import contextlib
    import pathlib
    from openpyxl.writer.excel import ExcelWriter

    obj = tmp_path / "obj.yaml"
    obj.write_text("{}\n", encoding="utf-8")
    out = tmp_path / "form.xlsx"
    out.write_bytes(b"previous form")
    writers = []

    def partial_write(self):
        writers.append(self)
        self._archive.writestr("partial", b"x" * 64)
        raise OSError(28, "No space left on device")

    real_unlink = pathlib.Path.unlink

    def windows_locked_unlink(self, missing_ok=False):
        if self.name.endswith(".tmp"):
            raise PermissionError(32, "file in use by another process")
        return real_unlink(self, missing_ok=missing_ok)

    monkeypatch.setattr(ExcelWriter, "write_data", partial_write)
    monkeypatch.setattr(pathlib.Path, "unlink", windows_locked_unlink)
    try:
        result = _build(tmp_path, obj, obj, obj, out=out)
        assert result.exit_code == 1, result.output
        assert "No space left on device" in result.output
        assert out.read_bytes() == b"previous form"
    finally:
        for writer in writers:
            with contextlib.suppress(ValueError):
                writer._archive.close()


@pytest.mark.parametrize("kind", ["input", "unrelated", "hardlink", "symlink"])
def test_build_survey_form_preserves_preexisting_staging(tmp_path, kind):
    """#554: the old shared staging name belongs to its existing owner."""
    import os
    from openpyxl import load_workbook

    obj = tmp_path / "obj.yaml"
    obj.write_text("{}\n", encoding="utf-8")
    out = tmp_path / "form.xlsx"
    stage = tmp_path / "form.xlsx.tmp"
    original = b"site_id: INPUT\nsite_name: Input Site\n"
    if kind in {"hardlink", "symlink"}:
        source = tmp_path / "protected.yaml"
        source.write_bytes(original)
        if kind == "hardlink":
            os.link(source, stage)
        else:
            try:
                stage.symlink_to(source)
            except OSError as exc:
                pytest.skip(f"symlink creation unavailable: {exc}")
    else:
        stage.write_bytes(original)

    result = _build(tmp_path, stage if kind == "input" else obj, obj, obj, out)
    assert result.exit_code == 0, result.output
    assert stage.read_bytes() == original
    if kind in {"hardlink", "symlink"}:
        assert os.path.samefile(stage, source)
    if kind == "symlink":
        assert stage.is_symlink()
    wb = load_workbook(out)
    try:
        assert wb.sheetnames == ["survey", "choices", "settings"]
        if kind == "input":
            assert wb["settings"].cell(2, 1).value == "Input Site Field Sampling"
    finally:
        wb.close()
    assert list(tmp_path.glob("*.tmp")) == [stage]


def test_build_survey_form_real_failed_save_removes_owned_stage(tmp_path, monkeypatch):
    """#551/#554: native Windows cleanup must close the failed ZIP handle."""
    import contextlib
    import os
    from openpyxl.writer.excel import ExcelWriter

    obj = tmp_path / "obj.yaml"
    obj.write_text("{}\n", encoding="utf-8")
    out = tmp_path / "form.xlsx"
    original = b"previous form"
    out.write_bytes(original)
    stage = tmp_path / "form.xlsx.tmp"
    protected = tmp_path / "protected.xlsx"
    protected.write_bytes(original)
    os.link(protected, stage)
    writers = []

    def partial_write(self):
        writers.append(self)
        self._archive.writestr("partial", b"x" * 64)
        raise OSError(28, "No space left on device")

    monkeypatch.setattr(ExcelWriter, "write_data", partial_write)
    try:
        result = _build(tmp_path, obj, obj, obj, out)
        assert result.exit_code == 1, result.output
        assert isinstance(result.exception, SystemExit), result.exception
        assert "No space left on device" in result.output
        assert out.read_bytes() == original
        assert stage.read_bytes() == original
        assert os.path.samefile(protected, stage)
        assert sorted(p.name for p in tmp_path.iterdir()) == [
            "form.xlsx", "form.xlsx.tmp", "obj.yaml", "protected.xlsx"]
    finally:
        for writer in writers:
            with contextlib.suppress(ValueError):
                writer._archive.close()


def test_build_survey_form_concurrent_processes_preserve_observed_publication(tmp_path):
    """#564: a later recheck observes the first publication and preserves it."""
    import json
    import subprocess
    import sys
    from openpyxl import load_workbook

    obj = tmp_path / "obj.yaml"
    obj.write_text("{}\n", encoding="utf-8")
    for role in ("A", "B"):
        (tmp_path / f"{role}.yaml").write_text(
            f"site_id: {role}\nsite_name: Site {role}\n", encoding="utf-8")
    script = r'''
import json, os, sys, time
from pathlib import Path
from unittest.mock import patch
from click.testing import CliRunner
from openpyxl import Workbook, load_workbook
from autogis.adapters.cli import autogis
role, folder = sys.argv[1], Path(sys.argv[2])
def wait(name):
    deadline = time.monotonic() + 15
    while not (folder / name).exists():
        if time.monotonic() >= deadline:
            raise RuntimeError("barrier timeout: " + name)
        time.sleep(0.01)
real_save = Workbook.save
real_replace = os.replace

def save(workbook, handle):
    real_save(workbook, handle)
    (folder / (role + "_stage")).write_text(str(handle.name))
    (folder / (role + "_saved")).touch()
    wait("B_saved" if role == "A" else "A_published")
def publish(source, target):
    real_replace(source, target)
    wb = load_workbook(target)
    try:
        title = wb["settings"].cell(2, 1).value
    finally:
        wb.close()
    (folder / (role + "_published.json")).write_text(json.dumps(title))
    (folder / (role + "_published")).touch()
if role == "B":
    wait("A_saved")
with patch.object(Workbook, "save", save), patch("autogis.adapters.cli.os.replace", publish):
    result = CliRunner().invoke(autogis, ["envmon", "build-survey-form",
        "--site", str(folder / (role + ".yaml")), "--event", str(folder / "obj.yaml"),
        "--analytes", str(folder / "obj.yaml"), "--out", str(folder / "form.xlsx")])
print(result.output)
sys.exit(result.exit_code)
'''
    processes = [subprocess.Popen(
        [sys.executable, "-c", script, role, str(tmp_path)],
        stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
        for role in ("A", "B")]
    results = [process.communicate(timeout=30) for process in processes]
    assert processes[0].returncode == 0, results[0]
    assert "XLSForm written" in results[0][0]
    assert json.loads((tmp_path / "A_published.json").read_text()) == "Site A Field Sampling"
    assert processes[1].returncode == 1, results[1]
    assert "output changed during build" in results[1][0]
    assert "XLSForm written" not in results[1][0]
    assert not (tmp_path / "B_published.json").exists()
    assert (tmp_path / "A_stage").read_text() != (tmp_path / "B_stage").read_text()
    wb = load_workbook(tmp_path / "form.xlsx")
    try:
        assert wb["settings"].cell(2, 1).value == "Site A Field Sampling"
    finally:
        wb.close()
    assert list(tmp_path.glob("*.tmp")) == []


@pytest.mark.parametrize("umask,existing_mode,expected_mode", [
    (0o022, None, 0o644), (0o077, None, 0o600),
    (0o077, 0o640, 0o640), (0o022, 0o600, 0o600)])
def test_build_survey_form_preserves_posix_output_mode(
        tmp_path, umask, existing_mode, expected_mode):
    """#559: real CLI output honors umask or retains the destination mode."""
    import os
    import stat
    import subprocess
    import sys
    from openpyxl import load_workbook

    if os.name != "posix":
        pytest.skip("POSIX permission bits")
    obj = tmp_path / "obj.yaml"
    obj.write_text("{}\n", encoding="utf-8")
    out = tmp_path / "form.xlsx"
    if existing_mode is not None:
        out.write_bytes(b"previous form")
        out.chmod(existing_mode)
    script = r'''
import os, runpy, stat, sys
from openpyxl import Workbook
real_save = Workbook.save
def save(self, handle):
    print(f"STAGING_MODE={stat.S_IMODE(os.fstat(handle.fileno()).st_mode):04o}")
    return real_save(self, handle)
Workbook.save = save
sys.argv = ["autogis", *sys.argv[1:]]
runpy.run_module("autogis", run_name="__main__")
'''
    result = subprocess.run([
        sys.executable, "-c", script, "envmon", "build-survey-form",
        "--site", str(obj), "--event", str(obj), "--analytes", str(obj),
        "--out", str(out)], capture_output=True, text=True,
        preexec_fn=lambda: os.umask(umask), timeout=30)
    assert result.returncode == 0, result.stdout + result.stderr
    assert f"STAGING_MODE={expected_mode:04o}" in result.stdout
    assert stat.S_IMODE(out.stat().st_mode) == expected_mode
    wb = load_workbook(out)
    try:
        assert wb.sheetnames == ["survey", "choices", "settings"]
    finally:
        wb.close()
    assert list(tmp_path.glob("*.tmp")) == []


def test_build_survey_form_uuid_collision_preserves_foreign_stage(tmp_path, monkeypatch):
    """A failed exclusive open never gives us ownership of another file."""
    from types import SimpleNamespace
    from autogis.adapters import cli

    obj = tmp_path / "obj.yaml"
    obj.write_text("{}\n", encoding="utf-8")
    out = tmp_path / "form.xlsx"
    out.write_bytes(b"previous form")
    stage = tmp_path / "collision.tmp"
    stage.write_bytes(b"another owner's file")
    monkeypatch.setattr(cli.uuid, "uuid4", lambda: SimpleNamespace(hex="collision"))
    result = _build(tmp_path, obj, obj, obj, out)
    assert result.exit_code == 1, result.output
    assert "cannot write" in result.output
    assert stage.read_bytes() == b"another owner's file"
    assert out.read_bytes() == b"previous form"


@pytest.mark.parametrize("seam", ["fstat", "fchown", "fchmod"])
def test_build_survey_form_mode_error_preserves_existing_form(tmp_path, monkeypatch, seam):
    """Stage metadata failure cannot write content, publish or leak our stage."""
    import os
    from types import SimpleNamespace
    from openpyxl import Workbook
    from autogis.adapters import cli

    if os.name != "posix":
        pytest.skip("POSIX permission operations")
    obj = tmp_path / "obj.yaml"
    obj.write_text("{}\n", encoding="utf-8")
    out = tmp_path / "form.xlsx"
    out.write_bytes(b"previous form")
    before = out.lstat()
    if seam == "fchown":
        monkeypatch.setattr(cli.os, "fstat", lambda fd: SimpleNamespace(
            st_uid=before.st_uid + 1, st_gid=before.st_gid + 1))

    def denied(*args):
        raise PermissionError(f"{seam} denied")

    monkeypatch.setattr(cli.os, seam, denied)
    saves = []
    monkeypatch.setattr(Workbook, "save", lambda *args: saves.append(args))
    result = _build(tmp_path, obj, obj, obj, out)
    assert result.exit_code == 1, result.output
    assert f"{seam} denied" in result.output
    assert "XLSForm written" not in result.output
    assert saves == []
    assert out.read_bytes() == b"previous form"
    after = out.lstat()
    assert (after.st_uid, after.st_gid, after.st_mode, after.st_ino) == (
        before.st_uid, before.st_gid, before.st_mode, before.st_ino)
    assert list(tmp_path.glob("*.tmp")) == []


@pytest.mark.parametrize("change", ["hardlink", "symlink", "created", "deleted", "identity", "mode", "content"])
def test_build_survey_form_rechecks_changed_output(tmp_path, monkeypatch, change):
    """#564: changes observed after save cannot be silently replaced."""
    import os
    from openpyxl import Workbook

    obj = tmp_path / "obj.yaml"
    obj.write_text("{}\n", encoding="utf-8")
    out = tmp_path / "form.xlsx"
    if change != "created":
        out.write_bytes(b"previous form")
    protected = tmp_path / "protected.xlsx"
    protected.write_bytes(b"protected form")
    observed = []
    real_save = Workbook.save

    def snapshot():
        if not out.exists() and not out.is_symlink():
            return None
        data = out.lstat()
        return (data.st_dev, data.st_ino, data.st_mode, data.st_nlink,
                data.st_size, data.st_mtime_ns, data.st_ctime_ns,
                os.readlink(out) if out.is_symlink() else out.read_bytes())

    def mutate(workbook, handle):
        real_save(workbook, handle)
        if change == "hardlink":
            os.link(out, tmp_path / "consumer.xlsx")
        elif change == "symlink":
            out.unlink()
            out.symlink_to(protected)
        elif change == "created":
            out.write_bytes(b"created by another writer")
        elif change == "deleted":
            out.unlink()
        elif change == "identity":
            replacement = tmp_path / "replacement.xlsx"
            replacement.write_bytes(b"replacement by another writer")
            os.replace(replacement, out)
        elif change == "mode":
            out.chmod(0o444)
        else:
            out.write_bytes(b"updated by another writer")
        observed.append(snapshot())

    monkeypatch.setattr(Workbook, "save", mutate)
    try:
        result = _build(tmp_path, obj, obj, obj, out)
        assert result.exit_code in (1, 2), result.output
        assert "XLSForm written" not in result.output
        assert ("output changed during build" in result.output
                or "single-link regular file" in result.output)
        assert snapshot() == observed[0]
        assert protected.read_bytes() == b"protected form"
        if change == "hardlink":
            assert out.read_bytes() == (tmp_path / "consumer.xlsx").read_bytes()
            assert os.path.samefile(out, tmp_path / "consumer.xlsx")
        assert list(tmp_path.glob("*.tmp")) == []
    finally:
        if change == "mode":
            out.chmod(0o666)


@pytest.mark.parametrize("kind", ["ascii221", "ascii255", "unicode255"])
def test_build_survey_form_supports_native_basename_limits(tmp_path, monkeypatch, kind):
    """#565: a valid output basename must not grow into the staging basename."""
    import os
    from pathlib import Path
    from openpyxl import Workbook, load_workbook

    obj = tmp_path / "obj.yaml"
    obj.write_text("{}\n", encoding="utf-8")
    name = ("a" * (216 if kind == "ascii221" else 250) + ".xlsx")
    if kind == "unicode255":
        name = "é" * (125 if os.name == "posix" else 250) + ".xlsx"
    out = tmp_path / name
    if os.name == "nt":
        out = Path("\\\\?\\" + str(out.resolve()))
    stages = []
    real_save = Workbook.save

    def save(workbook, handle):
        stages.append(Path(handle.name))
        return real_save(workbook, handle)

    monkeypatch.setattr(Workbook, "save", save)
    try:
        result = _build(tmp_path, obj, obj, obj, out)
        assert result.exit_code == 0, result.output
        assert len(stages) == 1 and len(stages[0].name) <= 64
        wb = load_workbook(out)
        try:
            assert wb.sheetnames == ["survey", "choices", "settings"]
        finally:
            wb.close()
        assert list(tmp_path.glob("*.tmp")) == []
    finally:
        out.unlink(missing_ok=True)


@pytest.mark.parametrize("changed_during_build", [False, True])
def test_build_survey_form_initial_owner_and_mode_before_content(tmp_path, monkeypatch, changed_during_build):
    """#566: stage access comes from the accepted output, before workbook bytes."""
    import os
    import stat
    from types import SimpleNamespace
    from openpyxl import Workbook
    from autogis.adapters import cli
    from autogis.core.envmon import survey123_form_builder as builder

    if os.name != "posix":
        pytest.skip("POSIX ownership operations")
    obj = tmp_path / "obj.yaml"
    obj.write_text("{}\n", encoding="utf-8")
    out = tmp_path / "form.xlsx"
    out.write_bytes(b"previous form")
    out.chmod(0o640)
    initial = out.lstat()
    calls = []
    real_fstat, real_fchown, real_fchmod = os.fstat, os.fchown, os.fchmod
    real_save, real_build = Workbook.save, builder.build_xlsform

    def stage_owner(fd):
        calls.append("fstat")
        return SimpleNamespace(st_uid=initial.st_uid + 1, st_gid=initial.st_gid + 1)

    def copy_owner(fd, uid, gid):
        assert (uid, gid) == (initial.st_uid, initial.st_gid)
        calls.append("fchown")
        real_fchown(fd, uid, gid)

    def copy_mode(fd, mode):
        assert mode == stat.S_IMODE(initial.st_mode)
        calls.append("fchmod")
        real_fchmod(fd, mode)

    def build(*args):
        wb = real_build(*args)
        if changed_during_build:
            out.chmod(0o444)
        return wb

    def save(wb, handle):
        metadata = real_fstat(handle.fileno())
        assert (metadata.st_uid, metadata.st_gid, stat.S_IMODE(metadata.st_mode)) == (
            initial.st_uid, initial.st_gid, 0o640)
        assert calls == ["fstat", "fchown", "fchmod"]
        calls.append("save")
        return real_save(wb, handle)

    monkeypatch.setattr(cli.os, "fstat", stage_owner)
    monkeypatch.setattr(cli.os, "fchown", copy_owner)
    monkeypatch.setattr(cli.os, "fchmod", copy_mode)
    monkeypatch.setattr(builder, "build_xlsform", build)
    monkeypatch.setattr(Workbook, "save", save)
    try:
        result = _build(tmp_path, obj, obj, obj, out)
        assert result.exit_code == (1 if changed_during_build else 0), result.output
        assert calls == ["fstat", "fchown", "fchmod", "save"]
        if changed_during_build:
            assert "output changed during build" in result.output
            assert out.read_bytes() == b"previous form"
            assert stat.S_IMODE(out.lstat().st_mode) == 0o444
        assert list(tmp_path.glob("*.tmp")) == []
    finally:
        out.chmod(0o666)


def test_build_survey_form_same_owner_needs_no_chown(tmp_path, monkeypatch):
    import os
    from autogis.adapters import cli

    if os.name != "posix":
        pytest.skip("POSIX ownership operations")
    obj = tmp_path / "obj.yaml"
    obj.write_text("{}\n", encoding="utf-8")
    out = tmp_path / "form.xlsx"
    out.write_bytes(b"previous form")

    def unexpected_chown(*args):
        raise AssertionError("same-owner output should not require chown")

    monkeypatch.setattr(cli.os, "fchown", unexpected_chown)
    result = _build(tmp_path, obj, obj, obj, out)
    assert result.exit_code == 0, result.output
    assert "XLSForm written" in result.output
    assert list(tmp_path.glob("*.tmp")) == []


def test_build_survey_form_preserves_recreated_stage_after_publication(tmp_path, monkeypatch):
    """#567: successful replace consumes our stage; its pathname may have a new owner."""
    import os
    from pathlib import Path
    from openpyxl import load_workbook
    from autogis.adapters import cli

    obj = tmp_path / "obj.yaml"
    obj.write_text("{}\n", encoding="utf-8")
    out = tmp_path / "form.xlsx"
    out.write_bytes(b"previous form")
    foreign = b"new foreign stage owner"
    recreated = []
    real_replace = os.replace

    def publish_then_recreate(source, target):
        source = Path(source)
        consumed_inode = source.lstat().st_ino
        real_replace(source, target)
        source.write_bytes(foreign)
        metadata = source.lstat()
        assert metadata.st_ino != consumed_inode
        recreated.append((source, metadata.st_ino))

    monkeypatch.setattr(cli.os, "replace", publish_then_recreate)
    result = _build(tmp_path, obj, obj, obj, out)
    assert result.exit_code == 0, result.output
    assert "XLSForm written" in result.output
    wb = load_workbook(out)
    try:
        assert wb.sheetnames == ["survey", "choices", "settings"]
    finally:
        wb.close()
    assert len(recreated) == 1
    source, foreign_inode = recreated[0]
    assert source.exists(), "foreign stage pathname was removed after publication"
    assert source.read_bytes() == foreign and source.lstat().st_ino == foreign_inode
