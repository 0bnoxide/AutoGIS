"""reconcile-locations: unreadable inputs exit cleanly, a missing profile
sheet is not a PASS (probe finding)."""
import openpyxl
import pytest
from click.testing import CliRunner

from autogis.adapters.cli import main
from autogis.core.common.config import ConfigError, load_config

PROFILE = ("profile_id: P\nsheets:\n  - sheet_name: S\n"
           "    data_type: METALS\n    data_start_row: 2\n    id_column: A\n")


@pytest.fixture
def inputs(tmp_path):
    (tmp_path / "site.yaml").write_text("site_id: X\n", encoding="utf-8")
    (tmp_path / "p.yaml").write_text(PROFILE, encoding="utf-8")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A2"] = "MW-1"
    wb.save(tmp_path / "wb.xlsx")
    (tmp_path / "w.csv").write_text("LocationID\nMW-1\n", encoding="utf-8")
    (tmp_path / "bad").write_bytes(b"\xff\xfe\x00bad")
    (tmp_path / "a dír").mkdir()
    (tmp_path / "shape.yaml").write_text(
        "profile_id: P\nsheets:\n  S: {sheet_name: S}\n", encoding="utf-8")
    return tmp_path


def _run(d, workbook="wb.xlsx", profile="p.yaml", wells="w.csv"):
    return CliRunner().invoke(main, [
        "envmon", "reconcile-locations", str(d / "site.yaml"),
        str(d / workbook), "--profile", str(d / profile),
        "--wells-csv", str(d / wells)])


def test_valid_inputs_pass(inputs):
    r = _run(inputs)
    assert r.exit_code == 0, r.output
    assert "1 matched" in r.output


@pytest.mark.parametrize("arg", [
    {"profile": "a dír"}, {"profile": "bad"}, {"profile": "shape.yaml"},
    {"workbook": "a dír"}, {"workbook": "bad"}, {"workbook": "p.yaml"},
    {"wells": "a dír"}, {"wells": "bad"},
])
def test_unreadable_input_is_clean_error(inputs, arg):
    r = _run(inputs, **arg)
    assert r.exit_code == 1, r.output
    assert isinstance(r.exception, SystemExit), repr(r.exception)
    assert "Error:" in r.output


def test_missing_profile_sheet_is_not_pass(inputs):
    wb = openpyxl.Workbook()
    wb.active.title = "Other"
    wb.save(inputs / "wb.xlsx")
    r = _run(inputs)
    assert r.exit_code != 0, r.output
    assert "expected_sheet_missing" in r.output


def test_load_config_directory_is_config_error(tmp_path):
    with pytest.raises(ConfigError, match="could not be read"):
        load_config(tmp_path)


@pytest.mark.parametrize("profile", ["a dír", "bad", "shape.yaml"])
@pytest.mark.parametrize("argv", [
    lambda d, p: ["parser-profile", p, str(d / "wb.xlsx")],
    lambda d, p: ["evaluate-rpd", str(d / "wb.xlsx"), p, "--site", "X"],
    lambda d, p: ["reconcile-survey123-lab", "--survey", str(d / "w.csv"),
                  "--edd", str(d / "w.csv"), "--edd-profile", p, "--site", "X"],
], ids=["parser-profile", "evaluate-rpd", "reconcile-survey123-lab"])
def test_sibling_profile_loads_are_clean_errors(inputs, argv, profile):
    r = CliRunner().invoke(main, ["envmon"] + argv(inputs, str(inputs / profile)))
    assert r.exit_code == 1, r.output
    assert isinstance(r.exception, SystemExit), repr(r.exception)
    assert "Error:" in r.output


def test_parser_profile_bad_workbook_is_clean_error(inputs):
    r = CliRunner().invoke(main, ["envmon", "parser-profile",
                                  str(inputs / "p.yaml"), str(inputs / "bad")])
    assert r.exit_code == 1, r.output
    assert isinstance(r.exception, SystemExit), repr(r.exception)
    assert "cannot open workbook" in r.output
